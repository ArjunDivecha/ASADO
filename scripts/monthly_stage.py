#!/usr/bin/env python3
"""Build monthly T2/P2P/GDELT/Bloomberg artifacts and warehouse in isolation.

Input: ASADO/Data/monthly/runs/<run>/{context,backup,raw}; canonical raw inputs
are read only while the parent holds producer/loop/repair/NN locks.
Output: that run's stage/ASADO/Data/, frozen feature snapshot, per-input/source
and exact table-preservation receipts. No canonical writes or provider retries.
"""
import argparse
import ast
import calendar
from datetime import datetime
import io
import json
import logging
from pathlib import Path
import shutil
import subprocess
import sys
from typing import Any, Dict, List, Tuple
import numpy as np
import pandas as pd
from monthly_readiness import read, sha, atomic, freshness, PACIFIC, shift
from duckdb_lock_guard import guarded_connect

KEYS = ["date", "country", "variable"]
COLS = ["date", "country", "value", "variable", "source"]
RETURNS = {f"{n}MRet" for n in (1, 3, 6, 9, 12)}


def definitions(path, names, ns=None):
    """Load pinned pure calculation functions, never collector module/main."""
    ns = ns or dict(pd=pd, np=np, logger=logging.getLogger("monthly_stage"), Any=Any, Dict=Dict, List=List, Tuple=Tuple)
    tree = ast.parse(Path(path).read_text())
    functions = [n for n in tree.body if isinstance(n, ast.FunctionDef) and n.name in names]
    if {n.name for n in functions} != set(names):
        raise RuntimeError("Pinned pure calculation API changed")
    for n in tree.body:
        if isinstance(n, ast.Assign):
            try:
                value = ast.literal_eval(n.value)
            except (ValueError, TypeError):
                continue
            for target in n.targets:
                if isinstance(target, ast.Name):
                    ns[target.id] = value
    module = ast.Module(body=[ast.ImportFrom(module="__future__", names=[ast.alias(name="annotations")], level=0)] + functions, type_ignores=[])
    exec(compile(ast.fix_missing_locations(module), str(path), "exec"), ns)
    return ns


def merge_tidy(old, new, target, *, allow_returns=False):
    """Preserve every prior observed value. Only append current inputs or fill
    previously NULL forward-return cells when their window has just completed.
    Never revise a historical input based on today's provider vintage.
    """
    old, new = old.copy(), new.copy()
    for frame in (old, new):
        frame["date"] = pd.to_datetime(frame.date)
        if frame.duplicated(KEYS).any():
            raise RuntimeError("Duplicate source/country/date/variable keys")
    target = pd.Timestamp(target)
    new = new[(new.date == target) | (allow_returns & new.variable.isin(RETURNS) & (new.date < target) & (new.date >= target - pd.DateOffset(months=12)))].copy()
    old_index = old.set_index(KEYS)
    for row in new.itertuples(index=False):
        key = (row.date, row.country, row.variable)
        if key in old_index.index and pd.notna(old_index.loc[key, "value"]):
            if row.date == target and pd.notna(row.value) and not np.isclose(old_index.loc[key, "value"], row.value, rtol=1e-10, atol=1e-10):
                raise RuntimeError("Current input label already populated differently; reconcile before publishing: " + str(key))
            continue
        if row.date < target and row.variable not in RETURNS:
            continue
        old_index.loc[key, "value"] = row.value
    return old_index.reset_index().sort_values(KEYS).reset_index(drop=True)


def merge_wide(old, new, target, variable, *, observation=False):
    """Preserve old cells; append latest completed label, fill only return NULLs."""
    name = old.columns[0]
    old, new = old.copy(), new.copy()
    old[name] = pd.to_datetime(old[name]); new = new.rename(columns={new.columns[0]: name}); new[name] = pd.to_datetime(new[name])
    if old[name].duplicated().any() or new[name].duplicated().any():
        raise RuntimeError("Duplicate workbook dates: " + variable)
    old = old.set_index(name); new = new.set_index(name)
    if list(old.columns) != list(new.columns):
        raise RuntimeError("Workbook country order changed: " + variable)
    cutoff = pd.Timestamp(target)
    new = new[new.index == cutoff] if variable not in RETURNS else new[(new.index <= cutoff) & (new.index >= cutoff - pd.DateOffset(months=12))]
    for dt, row in new.iterrows():
        if dt not in old.index:
            old.loc[dt] = row
        elif variable in RETURNS:
            old.loc[dt] = old.loc[dt].where(old.loc[dt].notna(), row)
        else:
            overlap = old.loc[dt].notna() & row.notna()
            if overlap.any() and not np.allclose(old.loc[dt][overlap], row[overlap], rtol=1e-10, atol=1e-10):
                raise RuntimeError("Workbook current label changed: " + variable)
    return old.sort_index().reset_index()


def copy_inputs(c):
    root, run = Path(c["root"]), Path(c["run"])
    stage = run / "stage/ASADO"
    originals = read(c["originals"])
    for row in originals.values():
        p = Path(row["destination"])
        if row["existed"] and p.is_relative_to(root):
            dest = stage / p.relative_to(root)
            dest.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(row["backup"], dest)
    # Frozen historical GDELT cache/lookups needed to recompute causal rolling
    # features, never symlinked to live writable directories.
    end = c["observation_month"] + "-31"
    for rel in ("Data/gdelt/country_day", "Data/gdelt/manifests/country_day", "Data/gdelt/lookups"):
        for p in sorted((root / rel).glob("*")):
            if p.is_file() and (rel.endswith("lookups") or p.stem <= end):
                dest = stage / rel / p.name
                dest.parent.mkdir(parents=True, exist_ok=True)
                if not dest.exists():
                    shutil.copy2(p, dest)
    return stage


def replies(c):
    run = Path(c["run"])
    receipts = read(run / "collection_receipts.json")
    plan = read(run / "provider_manifest.json")
    if not receipts or {r["id"] for r in receipts} != {r["id"] for r in plan}:
        raise RuntimeError("Collection receipts incomplete")
    result = []
    for r in receipts:
        if sha(r["file"]) != r["sha256"]:
            raise RuntimeError("Raw provider receipt changed")
        item = read(r["file"])
        req = item["request"]
        if req != next(p for p in plan if p["id"] == r["id"]):
            raise RuntimeError("Raw request differs from manifest")
        batches = item["rows"].values() if req["kind"] == "t2" else [item["rows"]]
        for rows in batches:
            for row in rows:
                d = pd.Timestamp(row["date"])
                if not pd.Timestamp(req["start"]) <= d <= pd.Timestamp(req["end"]):
                    raise RuntimeError("Returned provider date outside approved range")
        result.append(item)
    return result


def t2_workbook(c, stage, raw):
    from openpyxl import load_workbook
    root = Path(c["root"])
    manifest = read(root / "scripts/config/t2_bbg_manifest.json")
    ns = definitions(root / "scripts/collect_t2_bloomberg.py", {"rows_to_monthly"})
    series = {}
    for item in raw:
        r = item["request"]
        if r["kind"] == "t2":
            for sec in r["securities"]:
                series[(sec, r["fields"][0], r["currency"])] = ns["rows_to_monthly"](item["rows"].get(sec, []), r["fields"][0])
    path = stage / "Data/work/t2/T2 Bloomberg Master.xlsx"
    wb = load_workbook(path)
    obs = pd.Period(c["observation_month"], freq="M")
    direct = {}
    warnings = []
    for s in manifest["sheets"]:
        name = s["sheet"]
        ws = wb[name]
        cols = s["columns"]
        meta = [(ws.cell(1,i+2).value, ws.cell(2,i+2).value) for i in range(len(cols))]
        if meta != [(v["ticker"], v["field"]) for v in cols]:
            raise RuntimeError("T2 raw metadata/country order differs from pinned manifest: " + name)
        currency = None if name == "10Yr Bond" else "USD"
        values = [series[(v["ticker"],v["field"],currency)].get(obs, np.nan) for v in cols]
        observed = [bool(pd.notna(v) and np.isfinite(v)) for v in values]
        if name in ("Tot Return Index ", "PX_LAST") and not all(observed):
            raise RuntimeError("Monthly reference source observations missing for " + name)
        # Preserve the legacy Fill=P only when this sheet has an actual new
        # observation. Never manufacture an annual/quarterly row from nothing.
        if not any(observed):
            warnings.append(dict(sheet=name, reason="No completed-month source observations; no row manufactured"))
            continue
        dates = {pd.Timestamp(ws.cell(i,1).value).to_period("M"):i for i in range(3,ws.max_row+1) if ws.cell(i,1).value}
        if obs in dates:
            old_values=[ws.cell(dates[obs],j).value for j in range(2,len(cols)+2)]
            if any(ok and (old is None or not np.isclose(old,v,rtol=1e-10,atol=1e-10)) for old,v,ok in zip(old_values,values,observed)):
                raise RuntimeError("T2 raw target month already present differently; reconcile")
            continue
        idx = ws.max_row + 1
        ws.cell(idx,1,obs.to_timestamp(how="end").normalize().to_pydatetime())
        for j, value in enumerate(values, 2):
            if not pd.notna(value):
                prior = next((ws.cell(i,j).value for i in range(idx-1,2,-1) if ws.cell(i,j).value is not None), None)
                value = prior
                warnings.append(dict(sheet=name, country_column=j, reason="Existing legacy Fill=P", last_observed_month=str(max(dates)) if dates else None))
            ws.cell(idx,j,None if value is None or not np.isfinite(float(value)) else float(value))
        direct[name] = values
    # Exact existing MCAP calculation, with no other raw-history rewrites.
    if "MCAP" in direct:
        raw_mcap = np.array([wb["MCAP"].cell(wb["MCAP"].max_row,j).value for j in range(2,36)],float)
        adj = raw_mcap * np.array(manifest["calc"]["mcap_split_factors"],float)
        for name, vals in (("MCAP Adj", adj), ("Mcap Weights", adj / np.nansum(adj))):
            ws=wb[name]; idx=ws.max_row+1;ws.cell(idx,1,obs.to_timestamp(how="end").normalize().to_pydatetime())
            for j,v in enumerate(vals,2):ws.cell(idx,j,None if not np.isfinite(v) else float(v))
    wb.save(path)
    wb.close()
    atomic(Path(c["run"]) / "T2_SOURCE_WARNINGS.json", warnings)


def p2p_and_t2(c, stage):
    import yfinance as yf
    import build_t2_master as builder
    import p2p_monthly as p2p
    import t2_normalize
    root, run = Path(c["root"]), Path(c["run"])
    work = stage / "Data/work/t2"
    yf.set_tz_cache_location(str(run / "yfinance_cache"))
    mapping = p2p.country_etfs()
    prior_meta=pd.read_excel(work/'P2P_Country_Historical_Scores.xlsx',sheet_name='_P2P_contract')
    prior_meta=dict(zip(prior_meta['key'],prior_meta['value']))
    use_cache=prior_meta.get('version')==p2p.VERSION and str(prior_meta.get('latest_complete_observation_month',''))>=c['observation_month']
    cached=pd.read_parquet(work/'P2P_Monthly_Adjusted_Prices.parquet') if use_cache else None
    series = {}
    for ticker in mapping.values():
        if cached is None:
            from monthly_owner import assert_owner
            assert_owner(c)
        data = cached[ticker] if cached is not None else yf.download(ticker, start="1999-01-01", end=c["month"]+"-01", interval="1mo", auto_adjust=True, progress=False, threads=False, timeout=30)["Close"]
        if isinstance(data,pd.DataFrame):data=data[ticker] if ticker in data else data.iloc[:,0]
        if data.empty:raise RuntimeError("P2P source empty: " + ticker)
        data.index=data.index.tz_localize(None)
        series[ticker]=data
        data.to_frame().to_parquet(run / ("p2p_"+ticker+".parquet"))
    prices=pd.DataFrame(series).sort_index()
    scores=p2p.compute_scores(prices,list(mapping.values()),as_of=pd.Timestamp(c["month"]+"-01"))
    old_scores=pd.read_excel(work/"P2P_Country_Historical_Scores.xlsx",sheet_name="Sheet1")
    contract=dict(scores.attrs["p2p_contract"])
    merged=merge_wide(old_scores,scores,c["month"]+"-01","P2P")
    merged.attrs["p2p_contract"]=contract
    prices.to_parquet(work/"P2P_Monthly_Adjusted_Prices.parquet")
    p2p.save_scores(merged,work/"P2P_Country_Historical_Scores.xlsx",work/"P2P_Monthly_Adjusted_Prices.parquet")
    builder.BLOOMBERG_FILE=work/"T2 Bloomberg Master.xlsx"
    builder.P2P_FILE=work/"P2P_Country_Historical_Scores.xlsx"
    candidate=builder.build_t2_master()  # In-memory only; never calls main/writers.
    previous=work/"T2 Master.xlsx"
    old=pd.read_excel(previous,sheet_name=None)
    new=pd.read_excel(io.BytesIO(candidate),sheet_name=None)
    if set(old)!=set(new):raise RuntimeError("T2 workbook schema changed")
    temp=previous.with_name("T2 Master.staged.xlsx")
    with pd.ExcelWriter(temp,engine="xlsxwriter") as writer:
        for name in old:
            merge_wide(old[name],new[name],c["month"]+"-01",name).to_excel(writer,sheet_name=name,index=False)
    temp.replace(previous)
    generated=run/"generated_t2";generated.mkdir()
    t2_normalize.normalize_data(previous,generated,write_xlsx=True)
    old_csv=pd.read_csv(work/"Normalized_T2_MasterCSV.csv",float_precision='round_trip')
    new_csv=pd.read_csv(generated/"Normalized_T2_MasterCSV.csv")
    merged_csv=merge_tidy(old_csv,new_csv,c["month"]+"-01",allow_returns=True)
    merged_csv.to_csv(work/"Normalized_T2_MasterCSV.csv",index=False)
    shutil.copy2(generated/"Normalized_T2_Master.xlsx",work/"Normalized_T2_Master.xlsx")
    for d in ("gdelt","econ"):
        dest=stage/"Data/work"/d;dest.mkdir(parents=True,exist_ok=True)
        shutil.copy2(previous,dest/previous.name);shutil.copy2(work/"Normalized_T2_MasterCSV.csv",dest/"Normalized_T2_MasterCSV.csv")


def gdelt(c, stage):
    root,run=Path(c["root"]),Path(c["run"])
    subprocess.run([sys.executable,str(root/"scripts/monthly_gdelt.py"),"--context",str(run/"context.json")],check=True)
    work=stage/"Data/work/gdelt"
    shutil.copy2(stage/"Data/gdelt/spreadsheet/GDELT.xlsx",work/"GDELT.xlsx")
    generated=run/"generated_gdelt";generated.mkdir()
    shutil.copy2(work/"GDELT.xlsx",generated/"GDELT.xlsx")
    shutil.copy2(work/"Normalized_T2_MasterCSV.csv",generated/"Normalized_T2_MasterCSV.csv")
    categories=Path(c["root"])/"Data/work/gdelt/Step Factor Categories.xlsx"
    if categories.exists():shutil.copy2(categories,generated/categories.name)
    import gdelt_normalize
    gdelt_normalize.normalize_gdelt(generated)
    old=pd.read_csv(work/"GDELT_Factors_MasterCSV.csv",float_precision='round_trip');new=pd.read_csv(generated/"GDELT_Factors_MasterCSV.csv")
    merge_tidy(old,new,c["month"]+"-01",allow_returns=True).to_csv(work/"GDELT_Factors_MasterCSV.csv",index=False)


def bloomberg(c, stage, raw):
    root,run=Path(c["root"]),Path(c["run"])
    ns=definitions(root/"scripts/collect_bloomberg.py",{"_coerce_float","to_monthly_first","_percentile_score","_merge_latest_market_cap","compute_passive_ratio_proxies","compute_passive_flow_distortion","compute_yield_curve_slope","compute_mipd","create_variable_catalog"})
    ns=definitions(root/"scripts/build_normalized_panel.py",{"_build_cross_sectional","_rolling_window_for_source","_build_time_series"},ns)
    base=pd.read_parquet(stage/"Data/processed/bloomberg_factors_panel.parquet");base.date=pd.to_datetime(base.date)
    families={};empty=[]
    for item in raw:
        req=item['request']
        if req['kind']!='bloomberg':continue
        role=req['roles'];frame=pd.DataFrame(item['rows']);family=role['family']
        if frame.empty:empty.append(req['securities'][0]);continue
        frame.date=ns['to_monthly_first'](pd.to_datetime(frame.date))
        for f in req['fields']:frame[f]=pd.to_numeric(frame.get(f),errors='coerce')
        if family=='country_etfs':
            frame=frame.groupby('date',as_index=False).agg(px_last=('PX_LAST','last'),shares_out=('EQY_SH_OUT','last'),fund_flow=('FUND_FLOW',lambda x:x.sum(min_count=1))).sort_values('date')
            parts=[]
            for var,kind in [('MS_Country_ETF_AUM_USD','aum'),('MS_Country_ETF_NetFlow_USD','flow'),('MS_ETF_NetCreation_Shares','creation')]:
                d=frame.copy();d['value']=d.px_last*d.shares_out if kind=='aum' else d.fund_flow if kind=='flow' else d.shares_out.diff();d['variable']=var;parts.append(d.dropna(subset=['value']))
            frame=pd.concat(parts)[['date','value','variable']]
        else:
            variable={'cds':'BBG_CDS_5Y','policy_rate_history':'BBG_WIRP_ImpliedRate','cpi_history':'BBG_ECFC_CPI','manufacturing_pmi':'BBG_PMI_Manufacturing'}.get(family)
            if family=='bonds':variable=next(v for v in role['outputs'] if v.startswith('BBG_Govt_Bond_'))
            frame=frame[['date',req['fields'][0]]].rename(columns={req['fields'][0]:'value'}).dropna(subset=['value']).groupby('date',as_index=False).last();frame['variable']=variable
        for country in role['countries']:
            part=frame.copy();part['country']=country;part['source']='bloomberg';families.setdefault(family,[]).append(part[COLS])
    required={'bonds','cds','policy_rate_history','cpi_history','manufacturing_pmi','country_etfs'}
    if set(families)!=required:raise RuntimeError('An entire Bloomberg family failed/empty: '+str(required-set(families)))
    fam={k:pd.concat(v,ignore_index=True) for k,v in families.items()}
    mapping=read(root/'config/country_mapping.json')['countries'];rawcap=stage/'Data/raw/bloomberg/worldbank_market_cap_usd.parquet'
    if not rawcap.exists():shutil.copy2(root/'Data/raw/bloomberg/worldbank_market_cap_usd.parquet',rawcap)
    caps=[]
    for row in pd.read_parquet(rawcap).to_dict('records'):
        if row.get('value') is None or not np.isfinite(float(row['value'])) or not str(row['date']).isdigit():continue
        for country,codes in mapping.items():
            if codes.get('wb')==row['countryiso3code'] and country in ns['T2_COUNTRIES']:caps.append(dict(date=pd.Timestamp(int(row['date']),12,1),country=country,market_cap_usd=float(row['value'])))
    ratios=ns['compute_passive_ratio_proxies'](fam['country_etfs'],pd.DataFrame(caps))
    parts=list(fam.values())+[ns['compute_yield_curve_slope'](fam['bonds']),ns['compute_mipd'](fam['cds']),ratios]
    parts.append(ns['compute_passive_flow_distortion'](pd.concat([fam['country_etfs'],ratios,base[base.variable.isin(['MS_Index_Weight','MS_Index_Weight_Change'])]],ignore_index=True)))
    fresh=pd.concat(parts,ignore_index=True);fresh=fresh[(fresh.date==pd.Timestamp(c['observation_month']+'-01'))&fresh.variable.isin(c['config']['bloomberg_base_variables'])][COLS].drop_duplicates(['date','country','variable','source'])
    if not np.isfinite(fresh.value).all():raise RuntimeError('Nonfinite Bloomberg additions')
    merged=merge_tidy(base.drop(columns='source'),fresh.drop(columns='source'),c['observation_month']+'-01');merged['source']='bloomberg';merged=merged[COLS]
    out=stage/'Data/processed';merged.to_parquet(out/'bloomberg_factors_panel.parquet',index=False);merged.to_csv(out/'bloomberg_factors_panel.csv',index=False)
    ns['create_variable_catalog'](merged).to_csv(out/'bloomberg_variable_catalog.csv',index=False)
    history=read(out/'bloomberg_run_history.json');history.append(dict(timestamp=datetime.now(PACIFIC).isoformat(),kind='bounded_monthly_nn_core',observation_month=c['observation_month'],request_manifest=str(run/'provider_manifest.json'),empty_responses=empty));atomic(out/'bloomberg_run_history.json',history)
    for name in ('bond_yields','cds_spreads','wirp','ecfc','pmi','country_etf_passive'):
        p=stage/f'Data/raw/bloomberg/{name}.parquet';old=pd.read_parquet(p);old.date=pd.to_datetime(old.date)
        add=fresh[fresh.variable.isin(old.variable.unique())]
        joined=merge_tidy(old.drop(columns='source'),add.drop(columns='source'),c['observation_month']+'-01');joined['source']='bloomberg';joined[COLS].to_parquet(p,index=False)
    selected=c['config']['bloomberg_selected_variables']
    normalized=pd.concat([ns['_build_cross_sectional'](merged),ns['_build_time_series'](merged)],ignore_index=True)
    normalized.date=pd.to_datetime(normalized.date)
    normalized=normalized[(normalized.date==pd.Timestamp(c['observation_month']+'-01'))&normalized.variable.isin(selected)]
    fresh.to_parquet(run/'new_bloomberg.parquet',index=False);normalized.to_parquet(run/'new_bloomberg_normalized.parquet',index=False)
    atomic(run/'BLOOMBERG_SOURCE_WARNINGS.json',dict(empty_responses=empty,market_cap_source_sha256=sha(rawcap),market_cap_latest_year=max((str(r['date']) for r in pd.read_parquet(rawcap).to_dict('records')),default=None),note='Cached annual denominator preserved; date of normalized output does not certify every underlying component has monthly updates'))


def warehouse(c, stage):
    root,run=Path(c['root']),Path(c['run']);cfg=c['config'];db=stage/'Data/asado.duckdb'
    con=guarded_connect(db,wait_budget_s=0)
    originals=read(c['originals']);original_db=originals[str(root/'Data/asado.duckdb')]['backup']
    changed={'t2_master','t2_raw','gdelt_panel','bloomberg_factors','normalized_panel'}
    try:
        con.execute("ATTACH '"+original_db.replace("'","''")+"' AS previous (READ_ONLY)")
        for table,path in [('t2_master',stage/'Data/work/t2/Normalized_T2_MasterCSV.csv'),('gdelt_panel',stage/'Data/work/gdelt/GDELT_Factors_MasterCSV.csv')]:
            frame=pd.read_csv(path,float_precision='round_trip');frame.date=pd.to_datetime(frame.date)
            candidates=frame[(frame.date==pd.Timestamp(c['month']+'-01'))&frame.value.notna()]
            if not set(candidates.country)<=set(cfg['countries']):
                raise RuntimeError('New monthly input country identity differs from NN contract')
            con.register('monthly_frame',frame)
            # Fill only formerly NULL targets; never rewrite an observed past value.
            con.execute(f'''UPDATE {table} AS old SET value=n.value FROM monthly_frame n
              WHERE old.date=CAST(n.date AS DATE) AND old.country=n.country AND old.variable=n.variable AND old.value IS NULL
              AND old.variable IN ('1MRet','3MRet','6MRet','9MRet','12MRet') AND n.value IS NOT NULL''')
            con.execute(f'''INSERT INTO {table} SELECT CAST(n.date AS DATE),n.country,n.value,n.variable FROM monthly_frame n
              WHERE (CAST(n.date AS DATE)=DATE '{c['month']}-01' OR
                (n.variable IN ('1MRet','3MRet','6MRet','9MRet','12MRet') AND CAST(n.date AS DATE)>=DATE '{shift(c['month'],-12)}-01' AND CAST(n.date AS DATE)<DATE '{c['month']}-01'))
              AND NOT EXISTS(SELECT 1 FROM {table} old WHERE old.date=CAST(n.date AS DATE) AND old.country=n.country AND old.variable=n.variable)''')
            con.unregister('monthly_frame')
        ns=definitions(root/'scripts/setup_duckdb.py',{'_canonical_variable_name'})
        raw=[]
        for name,frame in pd.read_excel(stage/'Data/work/t2/T2 Master.xlsx',sheet_name=None).items():
            if name in ns['T2_RETURN_SHEETS']:continue
            frame=frame.rename(columns={frame.columns[0]:'date'}).melt(id_vars='date',var_name='country',value_name='value').dropna(subset=['date','value']);frame['variable']=ns['_canonical_variable_name'](name);raw.append(frame)
        data=pd.concat(raw,ignore_index=True);data=data[pd.to_datetime(data.date)==pd.Timestamp(c['month']+'-01')]
        con.register('new_t2_raw',data)
        con.execute('''INSERT INTO t2_raw SELECT CAST(n.date AS DATE),n.country,n.value,n.variable FROM new_t2_raw n
          WHERE NOT EXISTS(SELECT 1 FROM t2_raw old WHERE old.date=CAST(n.date AS DATE) AND old.country=n.country AND old.variable=n.variable)''')
        con.unregister('new_t2_raw')
        for table,p in [('bloomberg_factors',run/'new_bloomberg.parquet'),('normalized_panel',run/'new_bloomberg_normalized.parquet')]:
            frame=pd.read_parquet(p);con.register('new_bbg',frame)
            cols='CAST(n.date AS DATE),n.country,n.value,n.variable,n.source'
            if table=='normalized_panel':cols+=',n.normalization,n.base_variable,n.normalization_origin,CAST(n.lookback_observations AS INTEGER),CAST(n.min_observations AS INTEGER)'
            con.execute(f'''INSERT INTO {table} SELECT {cols} FROM new_bbg n WHERE NOT EXISTS
              (SELECT 1 FROM {table} old WHERE old.date=CAST(n.date AS DATE) AND old.country=n.country AND old.variable=n.variable AND old.source=n.source)''')
            con.unregister('new_bbg')
        tables=[r[0] for r in con.execute("SELECT table_name FROM information_schema.tables WHERE table_catalog='asado' AND table_schema='main' AND table_type='BASE TABLE'").fetchall()]
        preservation={}
        for table in tables:
            # Exact multiset preservation of all original observations; NULL
            # forward targets may become realized under the approved window.
            where=" WHERE NOT(value IS NULL AND variable IN ('1MRet','3MRet','6MRet','9MRet','12MRet'))" if table in {'t2_master','gdelt_panel'} else ''
            lost=con.execute(f'SELECT count(*) FROM (SELECT * FROM previous."{table}"{where} EXCEPT ALL SELECT * FROM main."{table}")').fetchone()[0]
            if lost:raise RuntimeError('Original warehouse observations changed/lost: '+table)
            extra=con.execute(f'SELECT count(*) FROM (SELECT * FROM main."{table}" EXCEPT ALL SELECT * FROM previous."{table}")').fetchone()[0]
            if table not in changed and extra:raise RuntimeError('Unrelated warehouse table changed: '+table)
            preservation[table]=dict(original_observations_preserved=True,added_or_realized_rows=extra)
        counts=con.execute("""SELECT source,variable,strftime(date,'%Y-%m'),list(distinct country) FROM feature_panel_observed
          WHERE value IS NOT NULL AND isfinite(value) GROUP BY source,variable,date""").fetchall()
        fs=read(Path(c['nn_root'])/cfg['factor_set'])
        report=freshness(counts,fs['factors'],cfg['countries'],c['month'])
        fills=[r for r in read(run/'T2_SOURCE_WARNINGS.json',[]) if r['reason']=='Existing legacy Fill=P']
        if fills:
            report['warnings'].append(f'T2 upstream retained existing Fill=P for {len(fills)} cells; see {run/"T2_SOURCE_WARNINGS.json"}; these are carried source cells, not newly observed values')
        bbg_warnings=read(run/'BLOOMBERG_SOURCE_WARNINGS.json',{})
        if bbg_warnings.get('empty_responses'):
            report['warnings'].append(f'Bloomberg bounded source has {len(bbg_warnings["empty_responses"])} empty security replies; see {run/"BLOOMBERG_SOURCE_WARNINGS.json"}')
        if bbg_warnings.get('market_cap_latest_year'):
            report['warnings'].append('Bloomberg ETF ratio denominator uses cached annual World Bank market-cap year '+bbg_warnings['market_cap_latest_year']+'; source snapshot/hash retained, no monthly-denominator freshness claim')
        report['full_coverage']=report['passed'] and not report['warnings']
        atomic(run/'freshness_detailed.json',report)
        if not report['passed']:raise RuntimeError('Staged per-input/source/return gate failed: '+str(report['problems']))
        snap=run/'stage/feature_panel_observed.parquet'
        con.execute("COPY (SELECT * FROM feature_panel_observed) TO '"+str(snap).replace("'","''")+"' (FORMAT PARQUET)")
        atomic(run/'STAGED.json',dict(preservation=preservation,freshness=report,snapshot=str(snap),snapshot_sha256=sha(snap)))
    finally:
        con.close()


def main():
    ap=argparse.ArgumentParser(description=__doc__);ap.add_argument('--context',type=Path,required=True);args=ap.parse_args();c=read(args.context)
    if not read(Path(c['run'])/'COLLECTED.json'):raise RuntimeError('Verified collection required')
    stage=copy_inputs(c);raw=replies(c)
    t2_workbook(c,stage,raw);p2p_and_t2(c,stage);gdelt(c,stage);bloomberg(c,stage,raw);warehouse(c,stage)


if __name__=='__main__':main()
